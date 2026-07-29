# SPDX-FileCopyrightText: 2026 John-IDK
# SPDX-License-Identifier: MIT

from __future__ import annotations

import hashlib
import html
import ipaddress
import mimetypes
import os
import re
import subprocess
import tempfile
import time
import xml.etree.ElementTree as ET
from collections import deque
from collections.abc import Callable, Iterable
from email import policy
from email.parser import BytesParser
from pathlib import Path
from urllib.parse import urldefrag, urljoin, urlparse
from urllib.robotparser import RobotFileParser

import httpx
import httpcore
from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from .models import ParsedSection, SourceDocument
from .executables import ExecutableResolutionError, resolve_trusted_executable
from .network import (
    PublicHTTPTransport,
    is_public_address,
    resolve_public_addresses,
)


SUPPORTED_SUFFIXES = {".pdf", ".docx", ".xlsx", ".pptx", ".md", ".markdown", ".txt", ".eml"}
CRAWL_REDIRECT_STATUS_CODES = frozenset({301, 302, 303, 307, 308})
MAX_CRAWL_REDIRECTS = 10
MAX_ROBOTS_BYTES = 512_000


def _tesseract_executable(
    explicit: str | Path | None = None,
    *,
    untrusted_roots: Iterable[str | Path] = (),
) -> Path | None:
    try:
        return resolve_trusted_executable(
            "tesseract",
            explicit=explicit,
            untrusted_roots=(Path.cwd(), *untrusted_roots),
        )
    except ExecutableResolutionError:
        return None


def ocr_health(
    executable: str | Path | None = None,
    *,
    untrusted_roots: Iterable[str | Path] = (),
) -> dict[str, object]:
    resolved_executable = _tesseract_executable(
        executable,
        untrusted_roots=untrusted_roots,
    )
    if resolved_executable is None:
        return {"tesseract_available": False, "languages": [], "vi_available": False, "en_available": False}
    result = subprocess.run([str(resolved_executable), "--list-langs"], text=True, encoding="utf-8", errors="replace", capture_output=True)
    languages = [line.strip() for line in result.stdout.splitlines()[1:] if line.strip()] if result.returncode == 0 else []
    return {
        "tesseract_available": result.returncode == 0,
        "languages": languages,
        "vi_available": "vie" in languages,
        "en_available": "eng" in languages,
    }


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def detect_language(text: str) -> str:
    vietnamese = set("ăâđêôơưáàảãạấầẩẫậắằẳẵặéèẻẽẹếềểễệíìỉĩịóòỏõọốồổỗộớờởỡợúùủũụứừửữựýỳỷỹỵ")
    lowered = text.lower()
    if any(ch in lowered for ch in vietnamese):
        return "vi"
    folded_tokens = set(re.findall(r"[a-z]+", lowered))
    vietnamese_markers = {"chinh", "sach", "doanh", "thu", "khong", "nguoi", "cong", "ty", "duoc", "cua", "va", "hang", "tuan"}
    return "vi" if len(folded_tokens & vietnamese_markers) >= 2 else "en"


def chunk_sections(sections: list[ParsedSection], max_chars: int = 1800, overlap: int = 180) -> list[ParsedSection]:
    chunks: list[ParsedSection] = []
    for section in sections:
        text = section.text.strip()
        if not text:
            continue
        start = 0
        part = 1
        while start < len(text):
            end = min(len(text), start + max_chars)
            if end < len(text):
                boundary = max(text.rfind("\n", start, end), text.rfind(". ", start, end))
                if boundary > start + max_chars // 2:
                    end = boundary + 1
            chunk = text[start:end].strip()
            if chunk:
                locator = section.locator if len(text) <= max_chars else f"{section.locator}, part {part}"
                chunks.append(ParsedSection(locator, section.title, chunk))
            if end >= len(text):
                break
            start = max(start + 1, end - overlap)
            part += 1
    return chunks


def parse_path(path: str | Path) -> SourceDocument:
    source = Path(path).resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    raw = source.read_bytes()
    suffix = source.suffix.lower()
    if suffix in {".md", ".markdown", ".txt"}:
        sections = _parse_text(raw, suffix)
    elif suffix == ".pdf":
        sections = _parse_pdf(source)
    elif suffix == ".docx":
        sections = _parse_docx(source)
    elif suffix == ".xlsx":
        sections = _parse_xlsx(source)
    elif suffix == ".pptx":
        sections = _parse_pptx(source)
    elif suffix == ".eml":
        sections = _parse_eml(raw)
    else:
        raise ValueError(f"unsupported source type: {suffix}")
    combined = "\n".join(section.text for section in sections)
    return SourceDocument(
        source_uri=source.as_uri(),
        title=source.stem,
        media_type=mimetypes.guess_type(source.name)[0] or "application/octet-stream",
        checksum=sha256_bytes(raw),
        language=detect_language(combined),
        sections=sections,
    )


def iter_supported_files(folder: str | Path) -> list[Path]:
    root = Path(folder).resolve()
    if not root.is_dir():
        raise NotADirectoryError(root)
    paths: list[Path] = []
    for current, directories, files in os.walk(root, followlinks=False):
        current_path = Path(current)
        directories.sort()
        files.sort()
        for name in files:
            path = current_path / name
            if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES:
                paths.append(path)
    return sorted(paths)


def _parse_text(raw: bytes, suffix: str) -> list[ParsedSection]:
    text = raw.decode("utf-8-sig", errors="replace")
    sections: list[ParsedSection] = []
    current_title = "Document"
    current_lines: list[str] = []
    start_line = 1
    for line_number, line in enumerate(text.splitlines(), 1):
        if suffix in {".md", ".markdown"} and line.startswith("#"):
            if current_lines:
                sections.append(ParsedSection(f"lines {start_line}-{line_number - 1}", current_title, "\n".join(current_lines).strip()))
            current_title = line.lstrip("# ").strip() or "Section"
            current_lines = []
            start_line = line_number + 1
        else:
            current_lines.append(line)
    if current_lines or not sections:
        sections.append(ParsedSection(f"lines {start_line}-{max(start_line, len(text.splitlines()))}", current_title, "\n".join(current_lines).strip()))
    return [section for section in sections if section.text]


def _parse_pdf(path: Path) -> list[ParsedSection]:
    import pypdfium2 as pdfium

    sections: list[ParsedSection] = []
    executable = _tesseract_executable(untrusted_roots=(path.parent,))
    health = (
        ocr_health(executable, untrusted_roots=(path.parent,))
        if executable
        else {"languages": []}
    )
    ocr_languages = "+".join(language for language in ("vie", "eng") if language in health["languages"])
    document = pdfium.PdfDocument(str(path))
    try:
        for index in range(len(document)):
            page = document[index]
            try:
                text_page = page.get_textpage()
                try:
                    text = text_page.get_text_range().strip()
                finally:
                    text_page.close()
                if len(text) < 32 and executable and ocr_languages:
                    bitmap = page.render(scale=2, rotation=0)
                    try:
                        with tempfile.TemporaryDirectory() as temp_dir:
                            image_path = Path(temp_dir) / "page.png"
                            output_base = Path(temp_dir) / "ocr"
                            image = bitmap.to_pil()
                            try:
                                image.save(image_path, format="PNG")
                            finally:
                                image.close()
                            result = subprocess.run(
                                [
                                    executable,
                                    str(image_path),
                                    str(output_base),
                                    "-l",
                                    ocr_languages,
                                ],
                                check=False,
                                capture_output=True,
                            )
                            if (
                                result.returncode == 0
                                and output_base.with_suffix(".txt").exists()
                            ):
                                text = output_base.with_suffix(".txt").read_text(
                                    encoding="utf-8",
                                    errors="replace",
                                ).strip()
                    finally:
                        bitmap.close()
                if text:
                    sections.append(
                        ParsedSection(
                            f"page {index + 1}",
                            f"Page {index + 1}",
                            text,
                        )
                    )
            finally:
                page.close()
    finally:
        document.close()
    return sections


def _parse_docx(path: Path) -> list[ParsedSection]:
    document = Document(path)
    sections: list[ParsedSection] = []
    paragraph_text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
    if paragraph_text:
        sections.append(ParsedSection("document paragraphs", "Document", paragraph_text))
    for index, table in enumerate(document.tables, 1):
        rows = ["\t".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        if rows:
            sections.append(ParsedSection(f"table {index}", f"Table {index}", "\n".join(rows)))
    return sections


def _parse_xlsx(path: Path) -> list[ParsedSection]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sections: list[ParsedSection] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[str] = []
            max_row = 0
            max_column = 0
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    rows.append("\t".join(values))
                    max_row = row_index
                    max_column = max(max_column, len(values))
            if rows:
                sections.append(ParsedSection(f"sheet {sheet.title}!A1:{sheet.cell(max_row, max_column).coordinate}", sheet.title, "\n".join(rows)))
    finally:
        workbook.close()
    return sections


def _parse_pptx(path: Path) -> list[ParsedSection]:
    presentation = Presentation(path)
    sections: list[ParsedSection] = []
    for index, slide in enumerate(presentation.slides, 1):
        pieces = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
        try:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                pieces.append(f"Notes: {notes}")
        except (AttributeError, ValueError):
            pass
        if pieces:
            sections.append(ParsedSection(f"slide {index}", f"Slide {index}", "\n".join(pieces)))
    return sections


def _parse_eml(raw: bytes) -> list[ParsedSection]:
    message = BytesParser(policy=policy.default).parsebytes(raw)
    headers = "\n".join(f"{name}: {message.get(name, '')}" for name in ("Subject", "From", "To", "Date"))
    sections = [ParsedSection("message headers", "Headers", headers)]
    body_parts: list[str] = []
    attachments: list[str] = []
    for part in message.walk():
        disposition = part.get_content_disposition()
        if disposition == "attachment":
            filename = part.get_filename() or "unnamed attachment"
            payload = part.get_payload(decode=True) or b""
            suffix = Path(filename).suffix.lower()
            if payload and suffix in SUPPORTED_SUFFIXES:
                with tempfile.TemporaryDirectory() as temp_dir:
                    attachment_path = Path(temp_dir) / Path(filename).name
                    attachment_path.write_bytes(payload)
                    try:
                        parsed = parse_path(attachment_path)
                        for section in parsed.sections:
                            sections.append(ParsedSection(f"attachment {filename} / {section.locator}", f"{filename}: {section.title}", section.text))
                    except (ValueError, OSError):
                        attachments.append(filename)
            else:
                attachments.append(filename)
            continue
        if part.get_content_type() == "text/plain":
            body_parts.append(part.get_content())
        elif part.get_content_type() == "text/html" and not body_parts:
            soup = BeautifulSoup(part.get_content(), "html.parser")
            body_parts.append(soup.get_text("\n", strip=True))
    if body_parts:
        sections.append(ParsedSection("message body", "Body", "\n".join(body_parts)))
    if attachments:
        sections.append(ParsedSection("attachments", "Attachments", "\n".join(attachments)))
    return sections


def _clean_html_page(url: str, content: str) -> tuple[str, str, list[str]]:
    soup = BeautifulSoup(content, "html.parser")
    for element in soup(["script", "style", "nav", "footer", "noscript"]):
        element.decompose()
    title = soup.title.get_text(" ", strip=True) if soup.title else url
    main = soup.find("main") or soup.find("article") or soup.body or soup
    text = re.sub(r"\n{3,}", "\n\n", main.get_text("\n", strip=True))
    links = [urljoin(url, anchor.get("href")) for anchor in soup.find_all("a", href=True)]
    return title, html.unescape(text), links


def is_public_host(host: str) -> bool:
    try:
        resolve_public_addresses(host, 80)
    except httpcore.ConnectError:
        return False
    return True


def normalize_crawl_domains(domains: Iterable[str]) -> set[str]:
    normalized: set[str] = set()
    for raw_domain in domains:
        domain = str(raw_domain).strip().casefold().lstrip(".").rstrip(".")
        if not domain:
            continue
        try:
            address = ipaddress.ip_address(domain)
        except ValueError:
            try:
                ascii_domain = domain.encode("idna").decode("ascii")
            except UnicodeError as exc:
                raise ValueError(f"invalid crawl allowlist domain: {raw_domain}") from exc
            labels = ascii_domain.split(".")
            if (
                len(ascii_domain) > 253
                or any(
                    not label
                    or len(label) > 63
                    or label.startswith("-")
                    or label.endswith("-")
                    or re.fullmatch(r"[a-z0-9-]+", label) is None
                    for label in labels
                )
            ):
                raise ValueError(f"invalid crawl allowlist domain: {raw_domain}")
            normalized.add(ascii_domain)
        else:
            if not is_public_address(address):
                raise ValueError(
                    f"crawl allowlist cannot contain non-global address: {raw_domain}"
                )
            normalized.add(str(address))
    if not normalized:
        raise ValueError("website crawl requires at least one explicit allowed domain")
    return normalized


def _is_allowed_crawl_url(url: str, allowed_domains: set[str]) -> bool:
    try:
        parsed = urlparse(url)
        parsed.port
    except ValueError:
        return False
    host = (parsed.hostname or "").casefold().rstrip(".")
    if parsed.scheme not in {"http", "https"}:
        return False
    if parsed.username is not None or parsed.password is not None:
        return False
    try:
        host = str(ipaddress.ip_address(host))
    except ValueError:
        try:
            host = host.encode("idna").decode("ascii")
        except UnicodeError:
            return False
    if not any(host == domain or host.endswith(f".{domain}") for domain in allowed_domains):
        return False
    try:
        return is_public_address(ipaddress.ip_address(host))
    except ValueError:
        return bool(host) and host != "localhost" and not host.endswith(".localhost")


def _crawler_client(user_agent: str) -> httpx.Client:
    return httpx.Client(
        headers={"User-Agent": user_agent},
        follow_redirects=False,
        timeout=20,
        transport=PublicHTTPTransport(),
        trust_env=False,
    )


def _get_with_validated_redirects(
    client: httpx.Client,
    url: str,
    allowed_domains: set[str],
    *,
    stream: bool,
    before_request: Callable[[str], bool] | None = None,
) -> httpx.Response | None:
    current_url = urldefrag(url)[0]
    for redirect_count in range(MAX_CRAWL_REDIRECTS + 1):
        if not _is_allowed_crawl_url(current_url, allowed_domains):
            return None
        if before_request is not None and not before_request(current_url):
            return None
        client.cookies.clear()
        request = client.build_request("GET", current_url)
        for sensitive_header in (
            "cookie",
            "authorization",
            "proxy-authorization",
        ):
            request.headers.pop(sensitive_header, None)
        try:
            response = client.send(request, stream=stream, follow_redirects=False)
        except httpx.HTTPError:
            return None
        location = response.headers.get("location")
        if response.status_code not in CRAWL_REDIRECT_STATUS_CODES or not location:
            return response
        response_url = str(response.url)
        response.close()
        if redirect_count == MAX_CRAWL_REDIRECTS:
            return None
        current_url = urldefrag(urljoin(response_url, location))[0]
    return None


def crawl_public(
    seeds: list[str],
    allowed_domains: set[str],
    *,
    user_agent: str,
    max_depth: int,
    max_pages: int,
    delay_seconds: float,
    max_bytes: int = 5_000_000,
    checkpoint: Callable[[], None] | None = None,
) -> list[SourceDocument]:
    normalized_domains = normalize_crawl_domains(allowed_domains)
    queue = deque((urldefrag(seed)[0], 0) for seed in seeds)
    visited: set[str] = set()
    documents: list[SourceDocument] = []
    robots: dict[str, RobotFileParser] = {}
    with _crawler_client(user_agent) as client:
        def robots_allows(url: str) -> bool:
            parsed = urlparse(url)
            origin = f"{parsed.scheme}://{parsed.netloc}"
            if origin not in robots:
                parser = RobotFileParser(urljoin(origin, "/robots.txt"))
                try:
                    response = _get_with_validated_redirects(
                        client,
                        parser.url,
                        normalized_domains,
                        stream=True,
                    )
                    if response is None:
                        parser.parse(["User-agent: *", "Disallow: /"])
                    else:
                        try:
                            if response.status_code in {401, 403} or response.status_code >= 500:
                                parser.parse(["User-agent: *", "Disallow: /"])
                            elif response.status_code >= 400:
                                parser.parse([])
                            else:
                                robots_bytes = bytearray()
                                oversized = False
                                for block in response.iter_bytes():
                                    robots_bytes.extend(block)
                                    if len(robots_bytes) > min(max_bytes, MAX_ROBOTS_BYTES):
                                        oversized = True
                                        break
                                if oversized:
                                    parser.parse(["User-agent: *", "Disallow: /"])
                                else:
                                    parser.parse(
                                        bytes(robots_bytes)
                                        .decode(response.encoding or "utf-8", errors="replace")
                                        .splitlines()
                                    )
                        finally:
                            response.close()
                except httpx.HTTPError:
                    parser.parse(["User-agent: *", "Disallow: /"])
                robots[origin] = parser
            return robots[origin].can_fetch(user_agent, url)

        while queue and len(documents) < max_pages and len(visited) < max_pages * 10:
            if checkpoint:
                checkpoint()
            url, depth = queue.popleft()
            if url in visited:
                continue
            visited.add(url)
            if not _is_allowed_crawl_url(url, normalized_domains):
                continue
            response = _get_with_validated_redirects(
                client,
                url,
                normalized_domains,
                stream=True,
                before_request=robots_allows,
            )
            if response is None:
                continue
            try:
                response.raise_for_status()
                final_url = str(response.url)
                final_parsed = urlparse(final_url)
                declared_length = int(response.headers.get("content-length", "0") or 0)
                if declared_length > max_bytes:
                    continue
                data = bytearray()
                oversized = False
                for block in response.iter_bytes():
                    data.extend(block)
                    if len(data) > max_bytes:
                        oversized = True
                        break
                if oversized:
                    continue
                raw = bytes(data)
                content_type = response.headers.get("content-type", "")
                encoding = response.encoding or "utf-8"
            finally:
                response.close()
            if checkpoint:
                checkpoint()
            if final_parsed.path.lower().endswith((".xml", ".xml.gz")) or "xml" in content_type:
                try:
                    root = ET.fromstring(raw)
                    locations = [element.text.strip() for element in root.iter() if element.tag.rsplit("}", 1)[-1] == "loc" and element.text]
                except ET.ParseError:
                    locations = []
                if depth <= max_depth:
                    for location in locations:
                        queue.append((urldefrag(location)[0], depth + 1))
                continue
            if "text/html" not in content_type:
                continue
            page_text = raw.decode(encoding, errors="replace")
            title, text, links = _clean_html_page(final_url, page_text)
            documents.append(SourceDocument(final_url, title, "text/html", sha256_bytes(raw), detect_language(text), [ParsedSection("web page", title, text)]))
            if depth < max_depth:
                for link in links:
                    queue.append((urldefrag(link)[0], depth + 1))
            if delay_seconds:
                time.sleep(delay_seconds)
    return documents
