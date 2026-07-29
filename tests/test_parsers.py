# SPDX-FileCopyrightText: 2026 John-IDK
# SPDX-License-Identifier: MIT

from __future__ import annotations

import os
import subprocess
import sys
import socket
from types import SimpleNamespace
from email.message import EmailMessage
from pathlib import Path

import httpcore
import httpx
import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

RUNTIME_SRC = (
    Path(__file__).resolve().parents[1]
    / "plugins"
    / "local-knowledge-suite"
    / "runtime"
    / "src"
)
if str(RUNTIME_SRC) not in sys.path:
    sys.path.insert(0, str(RUNTIME_SRC))

from local_knowledge.ckb import parsers  # noqa: E402
from local_knowledge.ckb import network as network_module  # noqa: E402
from local_knowledge.ckb.parsers import (  # noqa: E402
    chunk_sections,
    crawl_public,
    detect_language,
    is_public_host,
    iter_supported_files,
    parse_path,
)
from local_knowledge.ckb.network import (  # noqa: E402
    PublicHTTPTransport,
    PublicNetworkBackend,
)


def _write_text_pdf(path: Path, text: str) -> None:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(content)).encode("ascii")
        + b" >>\nstream\n"
        + content
        + b"\nendstream",
    ]
    payload = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for index, body in enumerate(objects, 1):
        offsets.append(len(payload))
        payload.extend(f"{index} 0 obj\n".encode("ascii"))
        payload.extend(body)
        payload.extend(b"\nendobj\n")
    xref = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    payload.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref}\n%%EOF\n"
        ).encode("ascii")
    )
    path.write_bytes(bytes(payload))


def test_parse_markdown_and_folder(tmp_path: Path) -> None:
    path = tmp_path / "policy.md"
    path.write_text("# Chính sách\nDoanh thu được báo cáo mỗi thứ Hai.\n\n## Owner\nFinance team owns the report.\n", encoding="utf-8")
    parsed = parse_path(path)
    assert parsed.language == "vi"
    assert len(parsed.sections) == 2
    assert iter_supported_files(tmp_path) == [path]
    assert detect_language("Chinh sach doanh thu hang tuan cua cong ty") == "vi"
    chunks = chunk_sections(parsed.sections, max_chars=24, overlap=4)
    assert len(chunks) > len(parsed.sections)
    assert all(chunk.locator for chunk in chunks)


def test_parse_docx_xlsx_pptx_pdf_and_eml(tmp_path: Path) -> None:
    docx = tmp_path / "brief.docx"
    document = Document()
    document.add_paragraph("Revenue policy belongs to Finance.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Finance"
    document.save(docx)

    xlsx = tmp_path / "metrics.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Metrics"
    sheet.append(["Metric", "Value"])
    sheet.append(["Revenue", 42])
    workbook.save(xlsx)

    pptx = tmp_path / "deck.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Quarterly report"
    slide.placeholders[1].text = "Revenue increased this quarter."
    presentation.save(pptx)

    pdf = tmp_path / "report.pdf"
    _write_text_pdf(pdf, "Revenue report is approved by Finance.")

    eml = tmp_path / "message.eml"
    message = EmailMessage()
    message["Subject"] = "Revenue reminder"
    message["From"] = "finance@example.com"
    message["To"] = "team@example.com"
    message.set_content("Submit the revenue report by Monday.")
    message.add_attachment(b"# Attached policy\nAttachment content is ingested.\n", maintype="text", subtype="markdown", filename="attached.md")
    eml.write_bytes(message.as_bytes())

    assert any("Revenue policy" in section.text for section in parse_path(docx).sections)
    assert any("Revenue\t42" in section.text for section in parse_path(xlsx).sections)
    assert any("Revenue increased" in section.text for section in parse_path(pptx).sections)
    assert any("approved by Finance" in section.text for section in parse_path(pdf).sections)
    assert any("Submit the revenue" in section.text for section in parse_path(eml).sections)
    assert any("Attachment content is ingested" in section.text for section in parse_path(eml).sections)


def test_pdf_scan_uses_pdfium_render_and_tesseract_ocr(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pdf = tmp_path / "scan.pdf"
    _write_text_pdf(pdf, "")
    monkeypatch.setattr(
        parsers,
        "_tesseract_executable",
        lambda *args, **kwargs: Path("C:/trusted/tesseract.exe"),
    )
    monkeypatch.setattr(
        parsers,
        "ocr_health",
        lambda *args, **kwargs: {
            "tesseract_available": True,
            "languages": ["eng"],
            "vi_available": False,
            "en_available": True,
        },
    )

    def fake_run(command: list[str], **_: object) -> SimpleNamespace:
        Path(command[2]).with_suffix(".txt").write_text(
            "OCR invoice approved by Finance.",
            encoding="utf-8",
        )
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr(subprocess, "run", fake_run)
    parsed = parse_path(pdf)
    assert parsed.sections[0].locator == "page 1"
    assert "OCR invoice approved" in parsed.sections[0].text


@pytest.mark.skipif(os.name != "nt", reason="Windows executable search hardening")
def test_ocr_health_ignores_project_local_and_relative_path_executables(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    project = tmp_path / "project"
    project.mkdir()
    (project / "tesseract.exe").write_bytes(b"project-controlled")
    monkeypatch.chdir(project)
    monkeypatch.setenv("PATH", f".{os.pathsep}{project}")

    def forbidden(*args: object, **kwargs: object) -> None:
        raise AssertionError("project-local tesseract must not execute")

    monkeypatch.setattr(subprocess, "run", forbidden)
    health = parsers.ocr_health()
    assert health["tesseract_available"] is False


@pytest.mark.skipif(os.name != "nt", reason="Windows executable search hardening")
def test_ocr_health_runs_absolute_trusted_path_candidate(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    project = tmp_path / "project"
    trusted = tmp_path / "trusted-bin"
    project.mkdir()
    trusted.mkdir()
    executable = trusted / "tesseract.exe"
    executable.write_bytes(b"trusted-test-double")
    monkeypatch.chdir(project)
    monkeypatch.setenv("PATH", str(trusted))
    calls: list[list[str]] = []

    def fake_run(command: list[str], **_: object) -> SimpleNamespace:
        calls.append(command)
        return SimpleNamespace(returncode=0, stdout="List of available languages:\neng\nvie\n")

    monkeypatch.setattr(subprocess, "run", fake_run)
    health = parsers.ocr_health()
    assert health["tesseract_available"] is True
    assert health["languages"] == ["eng", "vie"]
    assert calls == [[str(executable.resolve()), "--list-langs"]]


def test_invalid_pdf_fails_cleanly(tmp_path: Path) -> None:
    pdf = tmp_path / "broken.pdf"
    pdf.write_bytes(b"not a pdf")
    with pytest.raises(Exception):
        parse_path(pdf)


def test_public_crawler_honors_allowlist_and_robots(monkeypatch) -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/robots.txt":
            return httpx.Response(
                200,
                text="User-agent: *\nDisallow: /private\n",
            )
        if request.url.path == "/private":
            return httpx.Response(
                200,
                headers={"Content-Type": "text/html"},
                text="<html><body>Private secret</body></html>",
            )
        if request.url.path == "/child":
            return httpx.Response(
                200,
                headers={"Content-Type": "text/html"},
                text="<html><title>Child</title><main>Revenue child policy.</main></html>",
            )
        return httpx.Response(
            200,
            headers={"Content-Type": "text/html"},
            text=(
                '<html><title>Home</title><main>Public revenue policy.'
                '<a href="/child">Child</a><a href="/private">Private</a>'
                "</main></html>"
            ),
        )

    _install_mock_crawler_client(monkeypatch, handler)
    documents = crawl_public(
        ["https://public.example/"],
        {"public.example"},
        user_agent="CKB-Test",
        max_depth=1,
        max_pages=10,
        delay_seconds=0,
    )
    assert {document.title for document in documents} == {"Home", "Child"}
    assert all("Private secret" not in section.text for document in documents for section in document.sections)


def test_public_crawler_accepts_sitemap_seed(monkeypatch) -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/robots.txt":
            return httpx.Response(200, text="User-agent: *\nAllow: /\n")
        if request.url.path == "/sitemap.xml":
            return httpx.Response(
                200,
                headers={"Content-Type": "application/xml"},
                text=(
                    '<?xml version="1.0"?><urlset '
                    'xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">'
                    "<url><loc>https://public.example/child</loc></url></urlset>"
                ),
            )
        return httpx.Response(
            200,
            headers={"Content-Type": "text/html"},
            text="<html><title>Child</title><main>Child page.</main></html>",
        )

    _install_mock_crawler_client(monkeypatch, handler)
    documents = crawl_public(
        ["https://public.example/sitemap.xml"],
        {"public.example"},
        user_agent="CKB-Test",
        max_depth=2,
        max_pages=10,
        delay_seconds=0,
    )
    assert [document.title for document in documents] == ["Child"]


def test_public_crawler_rejects_redirect_outside_allowlist(monkeypatch) -> None:
    requested: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        requested.append(str(request.url))
        if request.url.path == "/robots.txt":
            return httpx.Response(200, text="User-agent: *\nAllow: /\n")
        return httpx.Response(
            302,
            headers={"Location": "https://outside.example/child"},
        )

    _install_mock_crawler_client(monkeypatch, handler)
    documents = crawl_public(
        ["https://public.example/redirect-external"],
        {"public.example"},
        user_agent="CKB-Test",
        max_depth=1,
        max_pages=10,
        delay_seconds=0,
    )
    assert documents == []
    assert requested == [
        "https://public.example/robots.txt",
        "https://public.example/redirect-external",
    ]


def test_public_crawler_requires_explicit_allowlist(monkeypatch) -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        raise AssertionError(f"request made without explicit allowlist: {request.url}")

    _install_mock_crawler_client(monkeypatch, handler)
    with pytest.raises(ValueError, match="explicit allowed domain"):
        crawl_public(
            ["https://public.example/"],
            set(),
            user_agent="CKB-Test",
            max_depth=0,
            max_pages=1,
            delay_seconds=0,
        )


@pytest.mark.parametrize(
    "allowlist",
    [
        {"https://public.example"},
        {"bad_domain.example"},
        {"127.0.0.1"},
        {"169.254.169.254"},
    ],
)
def test_public_crawler_rejects_malformed_or_non_public_allowlist(
    monkeypatch,
    allowlist: set[str],
) -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        raise AssertionError(f"request made with invalid allowlist: {request.url}")

    _install_mock_crawler_client(monkeypatch, handler)
    with pytest.raises(ValueError, match="allowlist|non-global"):
        crawl_public(
            ["https://public.example/"],
            allowlist,
            user_agent="CKB-Test",
            max_depth=0,
            max_pages=1,
            delay_seconds=0,
        )


def test_public_crawler_never_replays_response_cookies(monkeypatch) -> None:
    requested: list[tuple[str, str | None, str | None, str | None]] = []

    def handler(request: httpx.Request) -> httpx.Response:
        requested.append(
            (
                request.url.path,
                request.headers.get("cookie"),
                request.headers.get("authorization"),
                request.headers.get("proxy-authorization"),
            )
        )
        if request.url.path == "/robots.txt":
            return httpx.Response(
                200,
                headers={"Set-Cookie": "robots-secret=1"},
                text="User-agent: *\nAllow: /\n",
            )
        if request.url.path == "/":
            return httpx.Response(
                200,
                headers={
                    "Content-Type": "text/html",
                    "Set-Cookie": "page-secret=1",
                },
                text='<html><title>Home</title><main><a href="/child">Child</a></main></html>',
            )
        return httpx.Response(
            200,
            headers={"Content-Type": "text/html"},
            text="<html><title>Child</title><main>Child</main></html>",
        )

    client_class = httpx.Client
    transport = httpx.MockTransport(handler)
    monkeypatch.setattr(
        parsers,
        "_crawler_client",
        lambda user_agent: client_class(
            transport=transport,
            headers={
                "User-Agent": user_agent,
                "Authorization": "Bearer must-not-leak",
                "Proxy-Authorization": "Basic must-not-leak",
            },
            trust_env=False,
        ),
    )
    documents = crawl_public(
        ["https://public.example/"],
        {"public.example"},
        user_agent="CKB-Test",
        max_depth=1,
        max_pages=2,
        delay_seconds=0,
    )

    assert [document.title for document in documents] == ["Home", "Child"]
    assert requested == [
        ("/robots.txt", None, None, None),
        ("/", None, None, None),
        ("/child", None, None, None),
    ]


def test_public_crawler_caps_robots_and_fails_closed(monkeypatch) -> None:
    requested: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        requested.append(request.url.path)
        if request.url.path == "/robots.txt":
            return httpx.Response(
                200,
                content=b"User-agent: *\nAllow: /\n"
                + b"#" * parsers.MAX_ROBOTS_BYTES,
            )
        raise AssertionError("page fetched after oversized robots.txt")

    _install_mock_crawler_client(monkeypatch, handler)
    documents = crawl_public(
        ["https://public.example/"],
        {"public.example"},
        user_agent="CKB-Test",
        max_depth=0,
        max_pages=1,
        delay_seconds=0,
        max_bytes=parsers.MAX_ROBOTS_BYTES * 2,
    )

    assert documents == []
    assert requested == ["/robots.txt"]


def _install_mock_crawler_client(monkeypatch, handler) -> None:
    client_class = httpx.Client
    transport = httpx.MockTransport(handler)
    monkeypatch.setattr(
        parsers,
        "_crawler_client",
        lambda user_agent: client_class(
            transport=transport,
            headers={"User-Agent": user_agent},
            follow_redirects=False,
            timeout=20,
            trust_env=False,
        ),
    )


def test_public_crawler_validates_page_redirect_before_request(monkeypatch) -> None:
    requested: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        requested.append(str(request.url))
        if request.url.host == "public.example" and request.url.path == "/robots.txt":
            return httpx.Response(200, text="User-agent: *\nAllow: /\n")
        if request.url.host == "public.example" and request.url.path == "/start":
            return httpx.Response(302, headers={"Location": "http://127.0.0.1/internal"})
        if request.url.host == "127.0.0.1":
            return httpx.Response(200, headers={"Content-Type": "text/html"}, text="Private service")
        raise AssertionError(f"unexpected request: {request.url}")

    _install_mock_crawler_client(monkeypatch, handler)
    documents = crawl_public(
        ["https://public.example/start"],
        {"public.example"},
        user_agent="CKB-Test",
        max_depth=0,
        max_pages=1,
        delay_seconds=0,
    )

    assert documents == []
    assert requested == ["https://public.example/robots.txt", "https://public.example/start"]


def test_public_crawler_fails_closed_on_private_robots_redirect(monkeypatch) -> None:
    requested: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        requested.append(str(request.url))
        if request.url.host == "public.example" and request.url.path == "/robots.txt":
            return httpx.Response(302, headers={"Location": "http://127.0.0.1/robots.txt"})
        if request.url.host == "public.example" and request.url.path == "/":
            return httpx.Response(200, headers={"Content-Type": "text/html"}, text="<html><title>Public</title><main>Public page.</main></html>")
        if request.url.host == "127.0.0.1":
            return httpx.Response(200, text="User-agent: *\nDisallow: /\n")
        raise AssertionError(f"unexpected request: {request.url}")

    _install_mock_crawler_client(monkeypatch, handler)
    documents = crawl_public(
        ["https://public.example/"],
        {"public.example"},
        user_agent="CKB-Test",
        max_depth=0,
        max_pages=1,
        delay_seconds=0,
    )

    assert documents == []
    assert requested == ["https://public.example/robots.txt"]


def test_public_crawler_follows_valid_relative_redirect(monkeypatch) -> None:
    requested: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        requested.append(str(request.url))
        if request.url.path == "/robots.txt":
            return httpx.Response(200, text="User-agent: *\nAllow: /\n")
        if request.url.path == "/start":
            return httpx.Response(301, headers={"Location": "/final#section"})
        if request.url.path == "/final":
            return httpx.Response(200, headers={"Content-Type": "text/html"}, text="<html><title>Final</title><main>Redirected page.</main></html>")
        raise AssertionError(f"unexpected request: {request.url}")

    _install_mock_crawler_client(monkeypatch, handler)
    documents = crawl_public(
        ["https://public.example/start"],
        {"public.example"},
        user_agent="CKB-Test",
        max_depth=0,
        max_pages=1,
        delay_seconds=0,
    )

    assert [(document.source_uri, document.title) for document in documents] == [("https://public.example/final", "Final")]
    assert requested == ["https://public.example/robots.txt", "https://public.example/start", "https://public.example/final"]


def test_public_crawler_checks_robots_before_same_origin_redirect_target(monkeypatch) -> None:
    requested: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        requested.append(str(request.url))
        if request.url.path == "/robots.txt":
            return httpx.Response(200, text="User-agent: *\nDisallow: /private\n")
        if request.url.path == "/start":
            return httpx.Response(302, headers={"Location": "/private"})
        if request.url.path == "/private":
            raise AssertionError("redirect target was fetched despite robots.txt")
        raise AssertionError(f"unexpected request: {request.url}")

    _install_mock_crawler_client(monkeypatch, handler)
    documents = crawl_public(
        ["https://public.example/start"],
        {"public.example"},
        user_agent="CKB-Test",
        max_depth=0,
        max_pages=1,
        delay_seconds=0,
    )

    assert documents == []
    assert requested == ["https://public.example/robots.txt", "https://public.example/start"]


def test_public_crawler_checks_robots_before_cross_origin_redirect_target(monkeypatch) -> None:
    requested: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        requested.append(str(request.url))
        if request.url.host == "source.example" and request.url.path == "/robots.txt":
            return httpx.Response(200, text="User-agent: *\nAllow: /\n")
        if request.url.host == "source.example" and request.url.path == "/start":
            return httpx.Response(302, headers={"Location": "https://target.example/private"})
        if request.url.host == "target.example" and request.url.path == "/robots.txt":
            return httpx.Response(200, text="User-agent: *\nDisallow: /private\n")
        if request.url.host == "target.example" and request.url.path == "/private":
            raise AssertionError("cross-origin redirect target was fetched before its robots policy")
        raise AssertionError(f"unexpected request: {request.url}")

    _install_mock_crawler_client(monkeypatch, handler)
    documents = crawl_public(
        ["https://source.example/start"],
        {"source.example", "target.example"},
        user_agent="CKB-Test",
        max_depth=0,
        max_pages=1,
        delay_seconds=0,
    )

    assert documents == []
    assert requested == [
        "https://source.example/robots.txt",
        "https://source.example/start",
        "https://target.example/robots.txt",
    ]


def test_public_crawler_bounds_redirect_loops(monkeypatch) -> None:
    requested: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        requested.append(str(request.url))
        if request.url.path == "/robots.txt":
            return httpx.Response(200, text="User-agent: *\nAllow: /\n")
        if request.url.path == "/one":
            return httpx.Response(302, headers={"Location": "/two"})
        if request.url.path == "/two":
            return httpx.Response(307, headers={"Location": "/one"})
        raise AssertionError(f"unexpected request: {request.url}")

    _install_mock_crawler_client(monkeypatch, handler)
    documents = crawl_public(
        ["https://public.example/one"],
        {"public.example"},
        user_agent="CKB-Test",
        max_depth=0,
        max_pages=1,
        delay_seconds=0,
    )

    assert documents == []
    page_requests = [url for url in requested if not url.endswith("/robots.txt")]
    assert len(page_requests) == parsers.MAX_CRAWL_REDIRECTS + 1


def _dns_records(addresses: list[str], port: int) -> list[tuple[object, ...]]:
    records: list[tuple[object, ...]] = []
    for address in addresses:
        if ":" in address:
            records.append(
                (
                    socket.AF_INET6,
                    socket.SOCK_STREAM,
                    socket.IPPROTO_TCP,
                    "",
                    (address, port, 0, 0),
                )
            )
        else:
            records.append(
                (
                    socket.AF_INET,
                    socket.SOCK_STREAM,
                    socket.IPPROTO_TCP,
                    "",
                    (address, port),
                )
            )
    return records


class ScriptedNetworkStream(httpcore.NetworkStream):
    def __init__(self, response: bytes = b"") -> None:
        self.response = bytearray(response)
        self.writes: list[bytes] = []
        self.server_hostnames: list[str | None] = []
        self.closed = False

    def read(self, max_bytes: int, timeout: float | None = None) -> bytes:
        block = bytes(self.response[:max_bytes])
        del self.response[:max_bytes]
        return block

    def write(self, buffer: bytes, timeout: float | None = None) -> None:
        self.writes.append(bytes(buffer))

    def close(self) -> None:
        self.closed = True

    def start_tls(
        self,
        ssl_context,
        server_hostname: str | None = None,
        timeout: float | None = None,
    ):
        self.server_hostnames.append(server_hostname)
        return self

    def get_extra_info(self, info: str):
        if info == "is_readable":
            return False
        return None


class RecordingNetworkBackend(httpcore.NetworkBackend):
    def __init__(self, responses: list[bytes] | None = None) -> None:
        self.responses = list(responses or [b""])
        self.calls: list[tuple[str, int]] = []
        self.streams: list[ScriptedNetworkStream] = []

    def connect_tcp(
        self,
        host: str,
        port: int,
        timeout: float | None = None,
        local_address: str | None = None,
        socket_options=None,
    ) -> httpcore.NetworkStream:
        self.calls.append((host, port))
        response = self.responses.pop(0) if self.responses else b""
        stream = ScriptedNetworkStream(response)
        self.streams.append(stream)
        return stream

    def connect_unix_socket(
        self,
        path: str,
        timeout: float | None = None,
        socket_options=None,
    ) -> httpcore.NetworkStream:
        raise AssertionError("public crawler must not open UNIX sockets")

    def sleep(self, seconds: float) -> None:
        return None


@pytest.mark.parametrize(
    "addresses",
    [
        ["10.0.0.4"],
        ["169.254.169.254"],
        ["93.184.216.34", "127.0.0.1"],
        ["2606:2800:220:1:248:1893:25c8:1946", "fe80::1"],
        ["224.0.0.251"],
        ["ff02::1"],
        ["fec0::1"],
        ["64:ff9b::7f00:1"],
    ],
)
def test_public_network_backend_rejects_private_or_mixed_dns(
    addresses: list[str],
) -> None:
    direct = RecordingNetworkBackend()
    resolver_calls = 0

    def resolver(host: str, port: int, **_: object):
        nonlocal resolver_calls
        resolver_calls += 1
        return _dns_records(addresses, port)

    backend = PublicNetworkBackend(resolver=resolver, backend=direct)
    with pytest.raises(httpcore.ConnectError, match="non-global|mixed"):
        backend.connect_tcp("mixed.example", 443)

    assert resolver_calls == 1
    assert direct.calls == []


def test_public_network_backend_pins_first_resolution_to_numeric_ip() -> None:
    direct = RecordingNetworkBackend()
    resolver_calls = 0

    def rebinding_resolver(host: str, port: int, **_: object):
        nonlocal resolver_calls
        resolver_calls += 1
        address = "93.184.216.34" if resolver_calls == 1 else "127.0.0.1"
        return _dns_records([address], port)

    backend = PublicNetworkBackend(
        resolver=rebinding_resolver,
        backend=direct,
    )
    backend.connect_tcp("rebind.example", 443)

    assert resolver_calls == 1
    assert direct.calls == [("93.184.216.34", 443)]


def test_public_transport_preserves_original_host_and_tls_sni() -> None:
    response = (
        b"HTTP/1.1 200 OK\r\nContent-Length: 2\r\n"
        b"Connection: close\r\n\r\nOK"
    )
    direct = RecordingNetworkBackend([response])

    def resolver(host: str, port: int, **_: object):
        return _dns_records(["93.184.216.34"], port)

    transport = PublicHTTPTransport(
        network_backend=PublicNetworkBackend(
            resolver=resolver,
            backend=direct,
        )
    )
    with httpx.Client(transport=transport, trust_env=False) as client:
        fetched = client.get("https://public.example/policy")

    assert fetched.text == "OK"
    assert direct.calls == [("93.184.216.34", 443)]
    request_bytes = b"".join(direct.streams[0].writes)
    assert b"Host: public.example\r\n" in request_bytes
    assert direct.streams[0].server_hostnames == ["public.example"]


def test_public_transport_maps_stream_read_and_close_errors_to_httpx() -> None:
    request = httpx.Request("GET", "https://public.example/")

    class FailingCoreStream:
        def __iter__(self):
            raise httpcore.ReadError("read failed")
            yield b""

        def close(self) -> None:
            raise httpcore.ReadError("close failed")

    stream = network_module._ResponseStream(FailingCoreStream(), request)
    with pytest.raises(httpx.ReadError, match="read failed"):
        list(stream)
    with pytest.raises(httpx.ReadError, match="close failed"):
        stream.close()


def test_public_crawler_revalidates_redirect_dns_and_ignores_proxy_env(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    responses = [
        (
            b"HTTP/1.1 200 OK\r\nContent-Length: 23\r\n"
            b"Connection: close\r\n\r\nUser-agent: *\nAllow: /\n"
        ),
        (
            b"HTTP/1.1 302 Found\r\nContent-Length: 0\r\n"
            b"Connection: close\r\n"
            b"Location: https://private.example/secret\r\n\r\n"
        ),
    ]
    direct = RecordingNetworkBackend(responses)
    resolved_hosts: list[str] = []

    def resolver(host: str, port: int, **_: object):
        resolved_hosts.append(host)
        address = "93.184.216.34" if host == "public.example" else "10.0.0.8"
        return _dns_records([address], port)

    backend = PublicNetworkBackend(resolver=resolver, backend=direct)
    monkeypatch.setattr(
        parsers,
        "PublicHTTPTransport",
        lambda: PublicHTTPTransport(network_backend=backend),
    )
    monkeypatch.setenv("HTTPS_PROXY", "http://127.0.0.1:1")
    monkeypatch.setenv("ALL_PROXY", "http://127.0.0.1:1")
    monkeypatch.setenv("NO_PROXY", "")

    documents = crawl_public(
        ["https://public.example/start"],
        {"public.example", "private.example"},
        user_agent="CKB-Test",
        max_depth=0,
        max_pages=1,
        delay_seconds=0,
    )

    assert documents == []
    assert direct.calls == [
        ("93.184.216.34", 443),
        ("93.184.216.34", 443),
    ]
    assert resolved_hosts.count("public.example") == 2
    assert resolved_hosts.count("private.example") == 1


def test_private_hosts_are_not_public() -> None:
    assert not is_public_host("127.0.0.1")
    assert not is_public_host("localhost")
